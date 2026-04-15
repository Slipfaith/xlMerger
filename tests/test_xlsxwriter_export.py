# -*- coding: utf-8 -*-
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from utils.xlsxwriter_export import save_openpyxl_workbook_with_xlsxwriter


def test_save_openpyxl_workbook_with_xlsxwriter_preserves_formatting(tmp_path):
    src = tmp_path / "src.xlsx"
    out = tmp_path / "out.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "Header"
    ws["A1"].font = Font(bold=True, italic=True, color="FF0000")
    ws["A1"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws["A1"].number_format = "@"
    ws["A2"] = 123.45
    ws["A2"].number_format = "0.00"
    ws.column_dimensions["B"].width = 22
    ws.row_dimensions[2].height = 28
    ws["C1"] = "Merged"
    ws.merge_cells("C1:D1")
    ws.freeze_panes = "B2"
    wb.save(src)
    wb.close()

    wb_src = load_workbook(src)
    save_openpyxl_workbook_with_xlsxwriter(wb_src, str(out))
    wb_src.close()

    wb_out = load_workbook(out)
    ws_out = wb_out["Sheet1"]

    assert ws_out["A1"].value == "Header"
    assert ws_out["A1"].font.bold is True
    assert ws_out["A1"].font.italic is True
    assert ws_out["A1"].fill.fgColor.rgb is not None
    assert ws_out["A1"].fill.fgColor.rgb.upper().endswith("FFFF00")
    assert ws_out["A1"].alignment.horizontal == "center"
    assert ws_out["A1"].alignment.wrap_text is True
    assert ws_out["A2"].number_format == "0.00"
    assert ws_out["B1"].column_letter == "B"
    assert ws_out.column_dimensions["B"].width is not None
    # Excel/openpyxl normalize width units when re-reading the file.
    assert abs(ws_out.column_dimensions["B"].width - 22) < 1.0
    assert ws_out.row_dimensions[2].height == 28
    assert "C1:D1" in {str(rng) for rng in ws_out.merged_cells.ranges}
    assert ws_out.freeze_panes == "B2"

    wb_out.close()
